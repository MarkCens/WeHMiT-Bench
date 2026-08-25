# Import standard utilities for reproducible text normalization and splitting.
import argparse
import hashlib
import json
import os
import random
import re
import unicodedata
from pathlib import Path

# Read both worksheet values and cell styles from the reviewed Excel file.
import pandas as pd
from openpyxl import load_workbook


# Map workbook legend colors to explicit review-status fields.
DEFAULT_INPUT = "WeHMiT-Bench.xlsx"
DEFAULT_OUTPUT_DIR = "experiment_outputs"
RED_RGB = "FFFF0000"
YELLOW_RGB = "FFFFFF00"
BLUE_RGB = "FF0070C0"
GREEN_RGB = "FF00B050"


def parse_args():
    # Keep every preprocessing decision reproducible from command-line arguments.
    parser = argparse.ArgumentParser(description="Prepare WeHMiT-Bench from the expert-reviewed Excel workbook.")
    parser.add_argument("--input", default=DEFAULT_INPUT, help="Path to the reviewed Excel workbook.")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR, help="Directory for processed experiment files.")
    parser.add_argument("--sheet", default="Labels", help="Worksheet containing the reviewed records.")
    parser.add_argument("--split-manifest", default="", help="Optional stable id-to-split CSV; defaults to 01_split_manifest.csv in the output directory.")
    parser.add_argument("--seed", type=int, default=42, help="Random seed used for reproducible data splits.")
    parser.add_argument("--train-ratio", type=float, default=0.70, help="Training-set proportion.")
    parser.add_argument("--dev-ratio", type=float, default=0.15, help="Development-set proportion.")
    parser.add_argument("--force", action="store_true", help="Rebuild outputs even when the reviewed workbook is unchanged.")
    return parser.parse_args()


def color_rgb(color):
    # Return a stable empty value for theme colors and missing color objects.
    if color is None or color.type != "rgb" or color.rgb is None:
        return ""
    return str(color.rgb).upper()


def normalize_text(value):
    # Apply light Unicode and whitespace cleaning without rewriting title wording.
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = text.replace("\u200b", "").replace("\ufeff", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_title_key(value):
    # Remove layout punctuation only for duplicate matching, not model input.
    text = normalize_text(value).lower()
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[，。！？；：、,.!?;:'\"“”‘’《》【】\[\]()（）_\-—|丨]+", "", text)
    return text


def normalize_label(value):
    # Restrict the adjudicated task to the benchmark's two legal labels.
    label = normalize_text(value).lower()
    if label == "yes":
        return "Yes"
    if label == "no":
        return "No"
    raise ValueError(f"Unexpected label: {value}")


def resolve_column(sheet, aliases, required=True):
    # Resolve reviewed workbook fields by header so audited schema revisions remain explicit.
    headers = {
        normalize_text(sheet.cell(1, column_index).value): column_index
        for column_index in range(1, sheet.max_column + 1)
        if normalize_text(sheet.cell(1, column_index).value)
    }
    for alias in aliases:
        if alias in headers:
            return headers[alias]
    if required:
        raise ValueError(f"Missing required workbook column; expected one of: {aliases}")
    return None


def contains_any(text, phrases):
    # Phrase lists encode transparent review and title rules used below.
    return any(phrase in text for phrase in phrases)


def detect_review_flags(cells, communication_review, tcm_review):
    # Collect style and text evidence before interpreting individual flags.
    font_colors = [color_rgb(cell.font.color) for cell in cells]
    fill_colors = [color_rgb(cell.fill.fgColor) for cell in cells]
    review_text = f"{communication_review} {tcm_review}"

    # Convert the workbook color legend into explicit machine-readable flags.
    label_review_error = RED_RGB in font_colors and YELLOW_RGB in fill_colors
    out_of_domain_marked = BLUE_RGB in font_colors
    duplicate_marked = GREEN_RGB in font_colors or "重复" in review_text
    unavailable = contains_any(review_text, ["无法查看", "已被屏蔽", "文章已删除", "链接已不可访问", "公众号已迁移"])

    # Explanation-only issues are retained separately from label errors.
    explanation_issue = contains_any(review_text, ["解释理由不充分", "解释理由不够充分", "解释理由不准确", "解释理由错误", "判定理由错误"])
    return label_review_error, out_of_domain_marked, duplicate_marked, unavailable, explanation_issue


def detect_review_opinion(text):
    # Normalize recurring expert wording into correct and incorrect decisions.
    review_text = normalize_text(text)
    correct_phrases = ["根据标题判定准确", "根据标题判断准确", "根据标题的判定结果准确", "根据标题的判定理由和结果合理", "判定理由和结果合理", "判定结果和理由是对的"]
    wrong_phrases = ["根据标题判定错误", "根据标题判断错误", "根据标题判定结果错误", "根据标题判断结果错误", "根据标题判断失误", "根据标题判定失误", "AI判定结果错误", "AI判断结果错误"]

    # Explicit title-level statements receive more weight than article-body comments.
    says_correct = contains_any(review_text, correct_phrases)
    says_wrong = contains_any(review_text, wrong_phrases)
    if "标题确实属于健康误导信号" in review_text or "标题涉及严重的虚假" in review_text:
        says_correct = True
    if contains_any(review_text, ["标题确实无法直接判定", "标题的确没有呈现虚假", "标题本身不算 Misinformation"]):
        says_correct = True
    return says_correct, says_wrong


def derive_gold_label(source_label, label_review_error, communication_review, tcm_review):
    # Interpret both expert columns while preserving a traceable adjudication status.
    communication_correct, communication_wrong = detect_review_opinion(communication_review)
    tcm_correct, tcm_wrong = detect_review_opinion(tcm_review)

    # A contradictory communication review is not forced into a binary gold label.
    if communication_correct and communication_wrong:
        return source_label, "unresolved_communication_conflict", False

    # The communication scholar's title-level decision has priority in this benchmark.
    if communication_wrong:
        corrected_label = "No" if source_label == "Yes" else "Yes"
        return corrected_label, "flipped_by_communication_review", True
    if communication_correct:
        conflict = tcm_wrong and not tcm_correct
        status = "kept_title_review_over_content_conflict" if conflict else "kept_by_communication_review"
        return source_label, status, True

    # Use the TCM review when the communication review gives no title decision.
    if tcm_wrong:
        corrected_label = "No" if source_label == "Yes" else "Yes"
        return corrected_label, "flipped_by_tcm_review", True
    if tcm_correct:
        return source_label, "kept_by_tcm_review", True

    # A styled error without an explicit direction remains visible but is excluded from the core set.
    if label_review_error:
        return source_label, "unresolved_styled_label_error", False
    return source_label, "accepted_original_label", True


def infer_title_out_of_domain(title, communication_review, tcm_review, out_of_domain_marked):
    # Combine explicit title-level review with conservative semantic pattern checks.
    text = normalize_text(title)
    review_text = f"{communication_review} {tcm_review}"
    explicit_title_ood = contains_any(review_text, ["标题及内容与健康信息无关", "标题和内容与健康信息无关", "文章标题和内容，均与健康信息无关", "文章标题及内容均与健康信息无关", "标题及内容均不属于健康"])
    nonhealth_context = re.search(r"土方工程|施工阶段|结算阶段|审计要点|被刑拘|电视剧|电影特效|顶尖大律所|中方制裁|课外培训|学生家长|视频号|佛法|女性背包|爱宠|宠物|花千万别|婆婆带娃|育儿大法", text)
    health_signal = re.search(r"健康|养生|中医|药膳|食疗|食补|偏方|秘方|保健|营养|免疫|疾病|治病|癌|肿瘤|结节|血压|血糖|血脂|血栓|尿酸|痛风|肝病|肾病|胃病|心脏|肺|感冒|咳嗽|失眠|排毒|减肥|医生|用药|手术|化疗", text)

    # Blue styling alone is insufficient when the title clearly makes a health claim.
    if explicit_title_ood or nonhealth_context:
        return True
    if out_of_domain_marked and not health_signal:
        return True
    return False


def assign_topic(title, out_of_domain):
    # Assign exactly one broad topic for stratification and subgroup reporting.
    text = normalize_text(title)
    if out_of_domain:
        return "out_of_domain"

    # Apply specific disease groups before broad wellness and remedy categories.
    topic_rules = [
        ("cancer_tumor", r"癌|肿瘤|结节|乳腺|甲状腺"),
        ("cardiovascular", r"高血压|血压|血脂|血栓|血管|心梗|心脏|冠心"),
        ("diabetes_metabolic", r"糖尿病|血糖|胰岛素|痛风|尿酸|肥胖|减肥"),
        ("liver_kidney", r"肝病|肝炎|肝硬化|肾病|肾炎|尿毒症|肾虚"),

        # Keep organ-system and population groups in the same ordered rule table.
        ("gastrointestinal", r"胃病|胃炎|胃溃疡|肠道|幽门螺|便秘|腹泻"),
        ("respiratory_infection", r"感冒|咳嗽|咽喉|肺炎|感染|病毒|细菌"),
        ("women_children", r"妇科|女性|女人|儿童|孩子|宝宝|孕妇|更年期"),
        ("supplements_products", r"保健品|营养品|补充剂|益生菌|酵素|维生素|氢水|产品")
    ]
    for topic_name, pattern in topic_rules:
        if re.search(pattern, text):
            return topic_name

    # Broad traditional-health categories are used only after disease-specific checks.
    if re.search(r"偏方|秘方|祖传|验方|土方|良方", text):
        return "folk_remedy"
    if re.search(r"中医|养生|食疗|药膳|体质|经络|气血", text):
        return "traditional_wellness"
    return "general_health"


def assign_challenge_tags(title, topic, out_of_domain):
    text = normalize_text(title)
    tags = []

    # Challenge tags capture linguistic forms that commonly trigger model errors.
    if re.search(r"[?？]|吗|是否|真的|怎么|为何|为什么", text):
        tags.append("question_or_uncertainty")
    if re.search(r"辟谣|误区|别信|勿信|莫信|不靠谱|骗局|真相|揭秘|警惕|提醒", text):
        tags.append("debunking_or_warning")
    if re.search(r"[“”\"＂「」]|所谓|号称", text):
        tags.append("quoted_or_reported_claim")

    # Capture strong efficacy language and informal treatment mechanisms.
    if re.search(r"根治|断根|永不复发|神效|奇效|特效|百试百灵|一招见效|立竿见影|嗖嗖", text):
        tags.append("absolute_or_rapid_efficacy")
    if re.search(r"偏方|秘方|祖传|验方|土方|良方", text):
        tags.append("folk_remedy_keyword")
    if re.search(r"食疗|药膳|食补|吃.*治|喝.*治", text):
        tags.append("food_therapy_claim")

    # Track replacement, virality, and commercial persuasion separately.
    if re.search(r"不用吃药|不吃药|不用手术|不用化疗|替代|胜过药|比药", text):
        tags.append("replacement_of_standard_care")
    if re.search(r"医生不会告诉|专家不敢说|赶紧转|赶紧存|一定要看|千万别|看完吓|必须告诉", text):
        tags.append("viral_or_hidden_knowledge")
    if re.search(r"保健品|营养品|补充剂|产品|购买|价格", text):
        tags.append("commercial_product_claim")

    # Preserve domain collisions and otherwise untagged borderline wellness titles.
    if out_of_domain:
        tags.append("out_of_domain_keyword_collision")
    if topic in {"traditional_wellness", "general_health"} and not tags:
        tags.append("borderline_general_wellness")
    return "|".join(tags)


def load_reviewed_records(input_path, sheet_name):
    # Open the reviewed sheet without evaluating or altering workbook formulas.
    workbook = load_workbook(input_path, data_only=False)
    sheet = workbook[sheet_name]
    records = []

    # Support the released English schema while retaining compatibility with the reviewed source workbook.
    id_column = resolve_column(sheet, ["ID", "编号"])
    title_column = resolve_column(sheet, ["Title", "标题"])
    source_label_column = resolve_column(sheet, ["Preliminary Label (Generated by LLM)", "标签"])
    source_explanation_column = resolve_column(sheet, ["LLM Rationale", "解释"])
    final_label_column = resolve_column(sheet, ["Final Label (Verified by human experts)"], required=False)
    communication_review_column = resolve_column(sheet, ["Audit Rationale of Health Communication Scholar", "健康传播学者校验"])
    tcm_review_column = resolve_column(sheet, ["Audit Rationale of TCM Practitioner", "中医专家校验"])
    data_columns = [id_column, title_column, source_label_column, source_explanation_column, communication_review_column, tcm_review_column]
    if final_label_column is not None:
        data_columns.append(final_label_column)

    # Read values and styles together because review status is encoded by color.
    for row_index in range(2, sheet.max_row + 1):
        cells = [sheet.cell(row_index, column_index) for column_index in data_columns]
        source_id = normalize_text(sheet.cell(row_index, id_column).value)
        title = normalize_text(sheet.cell(row_index, title_column).value)
        if not source_id or not title:
            continue

        # Read the preliminary annotation, explicit verified label, and both audit-rationale columns.
        source_label = normalize_label(sheet.cell(row_index, source_label_column).value)
        source_explanation = normalize_text(sheet.cell(row_index, source_explanation_column).value)
        communication_review = normalize_text(sheet.cell(row_index, communication_review_column).value)
        tcm_review = normalize_text(sheet.cell(row_index, tcm_review_column).value)
        label_error, out_of_domain_marked, duplicate_marked, unavailable, explanation_issue = detect_review_flags(cells, communication_review, tcm_review)
        if final_label_column is not None:
            gold_label = normalize_label(sheet.cell(row_index, final_label_column).value)
            adjudication_status = "verified_final_label_unchanged" if gold_label == source_label else "verified_final_label_corrected"
            label_resolved = True
        else:
            gold_label, adjudication_status, label_resolved = derive_gold_label(source_label, label_error, communication_review, tcm_review)

        # Store both source and adjudicated labels for full auditability.
        record = {
            # Identity fields remain stable across every downstream experiment.
            "id": source_id,
            "source_row": row_index,
            "title": title,
            "title_key": normalize_title_key(title),
            # Annotation fields preserve the source and resolved benchmark target.
            "source_label": source_label,
            "gold_label": gold_label,
            "verified_final_label": gold_label if final_label_column is not None else "",
            "source_explanation": source_explanation,
            "health_communication_review": communication_review,
            "tcm_expert_review": tcm_review,
            "adjudication_status": adjudication_status,
            "label_resolved": label_resolved,
            # Review flags expose exclusions without discarding the source row.
            "style_label_error": label_error,
            "style_out_of_domain": out_of_domain_marked,
            "style_duplicate": duplicate_marked,
            "content_unavailable": unavailable,
            "explanation_issue": explanation_issue
        }
        records.append(record)
    return pd.DataFrame(records)


def add_dataset_fields(data_frame):
    # Compute normalized duplicate groups before applying any dataset exclusion.
    data_frame = data_frame.copy()
    first_id_by_title = data_frame.groupby("title_key", sort=False)["id"].transform("first")
    title_count = data_frame.groupby("title_key", sort=False)["id"].transform("count")

    # Combine reviewer duplicate marks with normalized-title duplicate detection.
    data_frame["duplicate_group_size"] = title_count.astype(int)
    data_frame["duplicate_of_id"] = data_frame["id"].where(data_frame["id"] != first_id_by_title, "")
    data_frame.loc[data_frame["id"] != first_id_by_title, "duplicate_of_id"] = first_id_by_title
    data_frame["is_duplicate"] = data_frame["style_duplicate"] | (data_frame["id"] != first_id_by_title)

    # Separate title-domain status from review colors that may describe article content.
    data_frame["title_out_of_domain"] = [infer_title_out_of_domain(title, communication_review, tcm_review, out_of_domain_marked) for title, communication_review, tcm_review, out_of_domain_marked in zip(data_frame["title"], data_frame["health_communication_review"], data_frame["tcm_expert_review"], data_frame["style_out_of_domain"])]
    data_frame["topic"] = [assign_topic(title, out_of_domain) for title, out_of_domain in zip(data_frame["title"], data_frame["title_out_of_domain"])]
    data_frame["challenge_tags"] = [assign_challenge_tags(title, topic, out_of_domain) for title, topic, out_of_domain in zip(data_frame["title"], data_frame["topic"], data_frame["title_out_of_domain"])]
    data_frame["is_challenge"] = data_frame["challenge_tags"].str.len() > 0
    return data_frame


def split_stratum(indices, train_ratio, dev_ratio, random_generator):
    # Shuffle each stratum independently with the shared seeded generator.
    shuffled_indices = list(indices)
    random_generator.shuffle(shuffled_indices)
    item_count = len(shuffled_indices)

    # Keep all three splits represented when a stratum has at least three items.
    if item_count == 1:
        return shuffled_indices, [], []
    if item_count == 2:
        return [shuffled_indices[0]], [], [shuffled_indices[1]]

    # Round requested proportions before reserving a nonempty test partition.
    train_count = max(1, int(round(item_count * train_ratio)))
    dev_count = max(1, int(round(item_count * dev_ratio)))
    if train_count + dev_count >= item_count:
        train_count = item_count - 2
        dev_count = 1
    return shuffled_indices[:train_count], shuffled_indices[train_count:train_count + dev_count], shuffled_indices[train_count + dev_count:]


def make_stratified_splits(core_data, train_ratio, dev_ratio, seed):
    # Build a deterministic assignment map keyed by original DataFrame indices.
    random_generator = random.Random(seed)
    split_map = {}
    stratify_key = core_data["gold_label"] + "__" + core_data["topic"]

    # Split within each label-topic stratum to preserve both dimensions.
    for group_name in sorted(stratify_key.unique()):
        group_indices = core_data.index[stratify_key == group_name].tolist()
        train_indices, dev_indices, test_indices = split_stratum(group_indices, train_ratio, dev_ratio, random_generator)

        # Commit all assignments only after the stratum has been split.
        split_map.update({index: "train" for index in train_indices})
        split_map.update({index: "dev" for index in dev_indices})
        split_map.update({index: "test" for index in test_indices})
    return pd.Series(split_map).reindex(core_data.index)


def assign_stable_splits(core_data, split_manifest_path, train_ratio, dev_ratio, seed):
    # Preserve published item-level partitions when labels are corrected in place.
    if split_manifest_path.exists():
        manifest = pd.read_csv(split_manifest_path, encoding="utf-8-sig", dtype={"id": str})
        required_columns = {"id", "split"}
        missing_columns = required_columns - set(manifest.columns)
        if missing_columns:
            raise ValueError(f"Split manifest is missing columns: {sorted(missing_columns)}")
        if manifest["id"].duplicated().any():
            raise ValueError("Split manifest contains duplicate ids.")
        invalid_splits = sorted(set(manifest["split"].astype(str)) - {"train", "dev", "test"})
        if invalid_splits:
            raise ValueError(f"Split manifest contains invalid split values: {invalid_splits}")
        split_map = dict(zip(manifest["id"].astype(str), manifest["split"].astype(str)))
        assignments = core_data["id"].astype(str).map(split_map)
    else:
        assignments = pd.Series(index=core_data.index, dtype="object")

    # Assign only genuinely new Core items, leaving every known id untouched.
    missing_mask = assignments.isna()
    if missing_mask.any():
        new_assignments = make_stratified_splits(core_data.loc[missing_mask], train_ratio, dev_ratio, seed)
        assignments.loc[missing_mask] = new_assignments
    return assignments


def write_data_frame(data_frame, output_path):
    # Replace a fully written UTF-8 CSV atomically to avoid partial local outputs.
    temporary_path = output_path.with_name(output_path.name + ".tmp")
    data_frame.to_csv(temporary_path, index=False, encoding="utf-8-sig")
    os.replace(temporary_path, output_path)


def write_json(value, output_path):
    # Use the same temporary-sibling pattern for summaries and completion markers.
    temporary_path = output_path.with_name(output_path.name + ".tmp")
    temporary_path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary_path, output_path)


def file_sha256(input_path):
    # Fingerprint the reviewed workbook so unchanged preprocessing can be reused.
    digest = hashlib.sha256()
    with input_path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def preprocessing_is_current(completion_path, output_dir, input_hash, split_manifest_hash, args):
    # Skip only when the source, split settings, and every required output match.
    if args.force or not completion_path.exists():
        return False
    completion = json.loads(completion_path.read_text(encoding="utf-8"))
    expected_files = ["01_all_reviewed_records.csv", "01_core_health_set.csv", "01_out_of_domain_set.csv", "01_challenge_set.csv", "01_train.csv", "01_dev.csv", "01_test.csv", "01_split_manifest.csv", "01_preprocessing_summary.json"]
    settings_match = completion.get("input_sha256") == input_hash and completion.get("split_manifest_sha256") == split_manifest_hash and completion.get("sheet") == args.sheet and completion.get("seed") == args.seed
    ratios_match = completion.get("train_ratio") == args.train_ratio and completion.get("dev_ratio") == args.dev_ratio
    return settings_match and ratios_match and all((output_dir / name).exists() for name in expected_files)


def main():
    # Resolve paths once and keep every generated artifact in one flat directory.
    args = parse_args()
    input_path = Path(args.input)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    completion_path = output_dir / "01_preprocessing_complete.json"
    split_manifest_path = Path(args.split_manifest) if args.split_manifest else output_dir / "01_split_manifest.csv"
    input_hash = file_sha256(input_path)
    split_manifest_hash = file_sha256(split_manifest_path) if split_manifest_path.exists() else ""

    # Reuse an intact deterministic preprocessing result on scheduled reruns.
    if preprocessing_is_current(completion_path, output_dir, input_hash, split_manifest_hash, args):
        print(completion_path.read_text(encoding="utf-8"))
        return

    # Load the reviewed workbook and create deterministic derived fields.
    all_data = load_reviewed_records(input_path, args.sheet)
    all_data = add_dataset_fields(all_data)
    all_data = all_data.sort_values("source_row").reset_index(drop=True)

    # Keep one resolved, accessible, in-domain title in the Core Health Set.
    core_mask = all_data["label_resolved"] & ~all_data["title_out_of_domain"] & ~all_data["is_duplicate"] & ~all_data["content_unavailable"]
    core_data = all_data.loc[core_mask].copy()
    core_data["split"] = assign_stable_splits(core_data, split_manifest_path, args.train_ratio, args.dev_ratio, args.seed)

    # The OOD set is deduplicated and retained separately for specificity testing.
    ood_mask = all_data["label_resolved"] & all_data["title_out_of_domain"] & ~all_data["is_duplicate"]
    ood_data = all_data.loc[ood_mask].copy()
    ood_data["split"] = "out_of_domain"
    challenge_data = core_data.loc[core_data["is_challenge"]].copy()

    # Write flat, machine-readable files for every downstream experiment.
    write_data_frame(all_data, output_dir / "01_all_reviewed_records.csv")
    write_data_frame(core_data, output_dir / "01_core_health_set.csv")
    write_data_frame(ood_data, output_dir / "01_out_of_domain_set.csv")
    write_data_frame(challenge_data, output_dir / "01_challenge_set.csv")
    write_data_frame(core_data[["id", "split"]].sort_values("id"), split_manifest_path)
    for split_name in ["train", "dev", "test"]:
        write_data_frame(core_data.loc[core_data["split"] == split_name], output_dir / f"01_{split_name}.csv")

    # Save a compact audit summary for the paper and reproducibility appendix.
    summary = {
        # Record corpus size and every exclusion category.
        "input_file": str(input_path),
        "split_manifest": str(split_manifest_path),
        "seed": args.seed,
        "total_rows": int(len(all_data)),
        "unique_title_keys": int(all_data["title_key"].nunique()),
        "core_health_rows": int(len(core_data)),
        "out_of_domain_rows": int(len(ood_data)),
        "duplicate_rows": int(all_data["is_duplicate"].sum()),
        "unavailable_rows": int(all_data["content_unavailable"].sum()),
        "unresolved_label_rows": int((~all_data["label_resolved"]).sum()),
        # Preserve label, split, topic, and adjudication distributions.
        "gold_label_counts": {key: int(value) for key, value in core_data["gold_label"].value_counts().to_dict().items()},
        "split_counts": {key: int(value) for key, value in core_data["split"].value_counts().to_dict().items()},
        "topic_counts": {key: int(value) for key, value in core_data["topic"].value_counts().to_dict().items()},
        "adjudication_counts": {key: int(value) for key, value in all_data["adjudication_status"].value_counts().to_dict().items()}
    }
    write_json(summary, output_dir / "01_preprocessing_summary.json")

    # Mark completion only after every expected CSV and summary has been replaced.
    split_manifest_hash = file_sha256(split_manifest_path)
    completion = {"status": "complete", "input_sha256": input_hash, "split_manifest_sha256": split_manifest_hash, "sheet": args.sheet, "seed": args.seed, "train_ratio": args.train_ratio, "dev_ratio": args.dev_ratio, "core_health_rows": int(len(core_data)), "out_of_domain_rows": int(len(ood_data))}
    write_json(completion, completion_path)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
